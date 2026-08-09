"""
Extracts raw text from whatever format the resume was uploaded in.
Kept separate from extractor.py deliberately -- text extraction is a
solved, boring problem and shouldn't be mixed with the LLM extraction
logic that actually needs iteration/testing.

Supported: .pdf, .docx, .txt
Anything else -> UnsupportedFormatError, caught by main.py and turned
into a user-facing Telegram message asking for a re-upload.
"""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class UnsupportedFormatError(Exception):
    pass


class EmptyResumeError(Exception):
    """Raised when extraction succeeds but yields near-zero text --
    usually a scanned/image-only PDF with no OCR layer."""
    pass


def load_text(file_path: str) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Resume file not found: {file_path}")

    suffix = path.suffix.lower()
    logger.info("Reading %s (%s)", path.name, suffix)
    if suffix == ".pdf":
        text = _load_pdf(path)
    elif suffix == ".docx":
        text = _load_docx(path)
    elif suffix == ".txt":
        text = path.read_text(errors="ignore")
    else:
        raise UnsupportedFormatError(
            f"'{suffix}' not supported. Please upload PDF, DOCX, or TXT."
        )

    text = text.strip()
    if len(text) < 50:
        # 50 chars is a generous floor -- even a sparse resume has more
        # than this. Below it, near-certainly a scanned image with no
        # text layer, and downstream LLM extraction would just hallucinate.
        raise EmptyResumeError(
            "Extracted almost no text. If this is a scanned resume, "
            "please upload a text-based PDF or DOCX instead."
        )
    logger.info("Extracted %d characters of text", len(text))
    return text


def _load_pdf(path: Path) -> str:
    import pdfplumber
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                chunks.append(page_text)
    return "\n".join(chunks)


def _load_docx(path: Path) -> str:
    import docx
    doc = docx.Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # tables (skills matrices, etc. are often in tables, not paragraphs)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n".join(paragraphs)